"""Excel test data for Robot Framework.

Robot Framework version of ``TestDataManager.getTestDataByRow()`` from the TypeScript Playwright
framework (ABM-India ``testdata/TestDataManager.ts``), with the same semantics:

- ``row`` is the physical, 1-based worksheet row; row 1 holds the headers, so row 2 is the first
  data row.
- ``sheet`` is a sheet name, a 0-based sheet index (pass an integer, e.g. ``${0}``), or omitted
  for the first sheet.
- Every non-empty header in row 1 becomes a key; empty cells become ``''``.
- Loaded workbooks are cached per file (reloaded when the file changes) and released with
  ``Clear Test Data Cache``, like ``TestDataManager.clearCache()``.

Differences, on purpose:

- Values of *secret columns* (default: ``Password``) are returned as Robot Framework ``Secret``
  objects, so they never appear in the log or report. Type them with Browser's ``Fill Secret``.
- ``required_columns`` fails clearly when a column is missing or empty in the row.
- Row 1 (the header row) cannot be requested as data.

Errors start with ``[TEST DATA]``, like the framework's other error categories.
"""

from __future__ import annotations

import os
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from robot.api import logger
from robot.api.deco import keyword, library
from robot.api.types import Secret


class TestDataError(Exception):
    ROBOT_SUPPRESS_NAME = True

    def __init__(self, message: str):
        super().__init__(f"[TEST DATA] {message}")


@library(scope="GLOBAL", auto_keywords=False)
class ExcelTestData:
    """Reads test data rows from ``.xlsx`` workbooks (header row + data rows)."""

    def __init__(self, secret_columns: str = "Password"):
        self._secret_columns = {name.strip().casefold() for name in secret_columns.split(",") if name.strip()}
        self._cache: dict[Path, tuple[int, Workbook]] = {}

    @keyword
    def get_test_data_by_row(
        self,
        workbook: str,
        row: int,
        sheet: int | str | None = None,
        *required_columns: str,
    ) -> dict[str, Any]:
        """Returns ``row`` of ``sheet`` in ``workbook`` as a dictionary keyed by the header row.

        ``workbook`` is a path, relative to the current working directory (the project root when
        run through the npm scripts) or absolute. ``required_columns`` must exist and be non-empty
        in the row. Secret columns are returned as ``Secret`` objects.

        | ${data}=    `Get Test Data By Row`    ${WORKBOOK}    2    LoginTestData    Username    Password
        | `Fill Text`      ${USERNAME_INPUT}    ${data}[Username]
        | `Fill Secret`    ${PASSWORD_INPUT}    ${data}[Password]
        """
        path = self._resolve(workbook)
        ws = self._sheet(self._load(path), sheet, path)
        if row < 2:
            raise TestDataError(
                f"Row {row} cannot be read as data from sheet '{ws.title}' of {path}: row 1 holds the headers, data starts at row 2."
            )
        headers = [cell.value for cell in ws[1]]
        if not any(header is not None and str(header) != "" for header in headers):
            raise TestDataError(f"Sheet '{ws.title}' of {path} has no header row (row 1 is empty).")
        values = [cell.value for cell in ws[row]] if row <= ws.max_row else []
        if all(value is None or value == "" for value in values):
            raise TestDataError(f"Row {row} not found or empty in sheet '{ws.title}' of {path} (last row: {ws.max_row}).")

        data: dict[str, Any] = {}
        for index, header in enumerate(headers):
            if header is None or str(header) == "":
                continue
            value = values[index] if index < len(values) and values[index] is not None else ""
            data[str(header)] = value

        for column in required_columns:
            if column not in data:
                raise TestDataError(
                    f"Column '{column}' not found in sheet '{ws.title}' of {path}. Available columns: {', '.join(data)}."
                )
            if str(data[column]).strip() == "":
                raise TestDataError(f"Column '{column}' is empty in row {row} of sheet '{ws.title}' of {path}.")

        for column in data:
            if column.casefold() in self._secret_columns and data[column] != "":
                data[column] = Secret(str(data[column]))
        logger.info(f"Read row {row} of sheet '{ws.title}' from {path.name}: columns {', '.join(data)}.")
        return data

    @keyword
    def clear_test_data_cache(self) -> None:
        """Releases all cached workbooks (use in suite teardown)."""
        count = len(self._cache)
        self._cache.clear()
        logger.info(f"Released {count} cached workbook(s).")

    # ------------------------------------------------------------------ helpers

    @staticmethod
    def _resolve(workbook: str) -> Path:
        if not workbook or not str(workbook).strip():
            raise TestDataError("No workbook path given. Set ORANGE_LOGIN_WORKBOOK (see .env.example).")
        path = Path(os.path.expandvars(str(workbook).strip())).expanduser()
        path = (path if path.is_absolute() else Path.cwd() / path).resolve()
        if not path.is_file():
            raise TestDataError(f"Excel file not found: {path}")
        return path

    def _load(self, path: Path) -> Workbook:
        mtime = path.stat().st_mtime_ns
        cached = self._cache.get(path)
        if cached and cached[0] == mtime:
            return cached[1]
        try:
            # The file is opened here and closed by the with-block, also when parsing fails;
            # a full (not read_only) load keeps nothing open afterwards.
            with path.open("rb") as handle:
                wb = load_workbook(handle, read_only=False, data_only=True)
        except Exception as error:
            raise TestDataError(f"Could not read Excel file {path}: {error}") from error
        self._cache[path] = (mtime, wb)
        logger.debug(f"Loaded workbook {path} (sheets: {', '.join(wb.sheetnames)}).")
        return wb

    @staticmethod
    def _sheet(wb: Workbook, sheet: int | str | None, path: Path) -> Worksheet:
        if sheet is None or sheet == "":
            if not wb.worksheets:
                raise TestDataError(f"Workbook {path} has no sheets.")
            return wb.worksheets[0]
        if isinstance(sheet, int):
            if not 0 <= sheet < len(wb.worksheets):
                raise TestDataError(f"Sheet index {sheet} not found in {path}. Available: 0-{len(wb.worksheets) - 1}.")
            return wb.worksheets[sheet]
        if sheet not in wb.sheetnames:
            raise TestDataError(f"Sheet '{sheet}' not found in {path}. Available: {', '.join(wb.sheetnames)}.")
        return wb[sheet]
